# -*- coding: utf-8 -*-
"""
monitor.py — Monitor de preços (demo de portfólio)

Fluxo:
  1. raspa o catálogo do site concorrente (aqui: books.toscrape.com, site
     público feito para treino de scraping)
  2. grava um snapshot diário em historico/precos_AAAA-MM-DD.csv
  3. compara com a sua tabela em dados_exemplo/meus_precos.csv
  4. gera saida/comparativo_precos.xlsx e docs/index.html

Uso:  python monitor.py

Autor: Marcos Galani — automação de processos
"""
from __future__ import annotations

import html as html_mod
import logging
import random
import re
import sys
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urljoin

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from scraper_base import baixar, buscar, criar_sessao, parsear, primeiro, texto

# ---------------------------------------------------------------------------
# Caminhos
# ---------------------------------------------------------------------------
RAIZ = Path(__file__).resolve().parent
ARQ_MEUS_PRECOS = RAIZ / "dados_exemplo" / "meus_precos.csv"
DIR_HISTORICO = RAIZ / "historico"
ARQ_XLSX = RAIZ / "saida" / "comparativo_precos.xlsx"
ARQ_HTML = RAIZ / "docs" / "index.html"

# CSVs no padrão brasileiro (abre direto no Excel): ponto e vírgula + vírgula decimal
CSV_OPCOES = dict(sep=";", decimal=",", encoding="utf-8-sig")

LIMIAR_ALERTA_PCT = 5.0   # acima disso "estou mais caro" vira alerta vermelho
MOEDA = "£"               # o site de treino publica preços em libras

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s",
                    datefmt="%H:%M:%S")
log = logging.getLogger("monitor")


# ===========================================================================
# FONTE: tudo que é específico do site raspado fica neste bloco.
# Para trocar de concorrente/marketplace, ajuste URL_PAGINA e os seletores
# dentro de extrair_produtos(). O resto do script não muda.
# ===========================================================================
URL_BASE = "https://books.toscrape.com/"
URL_PAGINA = "https://books.toscrape.com/catalogue/page-{n}.html"
PAGINAS = 3

# Seletores (tag, classe) usados em extrair_produtos()
SEL_CARD = ("article", "product_pod")        # um card por produto
SEL_TITULO = ("h3", None)                    # <h3><a title="..." href="...">
SEL_PRECO = ("p", "price_color")             # "£51.77"
SEL_DISPONIBILIDADE = ("p", "availability")  # "In stock"
SEL_ESTRELAS = ("p", "star-rating")          # classe extra: One..Five

_ESTRELAS = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5}


def converter_preco(txt: str) -> float | None:
    """'£51.77' -> 51.77 ; 'R$ 1.234,56' -> 1234.56 ; texto sem número -> None."""
    m = re.search(r"\d[\d.,]*", txt or "")
    if not m:
        return None
    s = m.group(0)
    if "," in s and "." in s:          # 1.234,56 -> 1234.56
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:                     # 51,77 -> 51.77
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def extrair_produtos(html: str, url_pagina: str) -> list[dict]:
    """Recebe o HTML de uma página do catálogo e devolve uma lista de produtos."""
    raiz = parsear(html)
    produtos = []
    for card in buscar(raiz, *SEL_CARD):
        h3 = primeiro(card, *SEL_TITULO)
        link = primeiro(h3, "a") if h3 else None
        if link is None:
            continue
        titulo = (link.attrs.get("title") or texto(link)).strip()
        url_produto = urljoin(url_pagina, link.attrs.get("href", ""))

        preco = converter_preco(texto(primeiro(card, *SEL_PRECO)))

        disp_txt = texto(primeiro(card, *SEL_DISPONIBILIDADE))
        disponibilidade = "Em estoque" if "in stock" in disp_txt.lower() else "Indisponível"

        estrelas = None
        p_est = primeiro(card, *SEL_ESTRELAS)
        if p_est is not None:
            for c in p_est.classes():
                if c.lower() in _ESTRELAS:
                    estrelas = _ESTRELAS[c.lower()]

        produtos.append({
            "titulo": titulo,
            "preco_concorrente": preco,
            "avaliacao_estrelas": estrelas,
            "disponibilidade": disponibilidade,
            "url": url_produto,
        })
    return produtos
# ======================= fim do bloco específico da fonte ===================


def raspar(paginas: int = PAGINAS) -> pd.DataFrame:
    sessao = criar_sessao()
    todos: list[dict] = []
    for n in range(1, paginas + 1):
        url = URL_PAGINA.format(n=n)
        log.info("baixando página %d/%d: %s", n, paginas, url)
        html = baixar(url, sessao)
        itens = extrair_produtos(html, url)
        log.info("  %d produtos encontrados", len(itens))
        todos.extend(itens)
    df = pd.DataFrame(todos)
    df.insert(0, "data", date.today().isoformat())
    df["coletado_em"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    df["fonte"] = URL_BASE
    return df


def salvar_historico(df: pd.DataFrame) -> Path:
    DIR_HISTORICO.mkdir(exist_ok=True)
    arq = DIR_HISTORICO / f"precos_{date.today().isoformat()}.csv"
    df.to_csv(arq, index=False, **CSV_OPCOES)
    log.info("histórico salvo: %s (%d linhas)", arq.name, len(df))
    return arq


def garantir_meus_precos(catalogo: pd.DataFrame) -> pd.DataFrame:
    """
    Carrega a tabela própria. Se ainda não existir (primeira execução da demo),
    gera uma tabela SINTÉTICA com ~20 títulos do catálogo e um "meu_preco"
    fictício — alguns acima, outros abaixo do concorrente — com seed fixa
    para o resultado ser reproduzível.
    """
    if ARQ_MEUS_PRECOS.exists():
        return pd.read_csv(ARQ_MEUS_PRECOS, **CSV_OPCOES)

    log.info("meus_precos.csv não existe; gerando tabela sintética de exemplo")
    rng = random.Random(42)
    amostra = catalogo.dropna(subset=["preco_concorrente"]).sample(n=20, random_state=42)
    linhas = []
    for i, (_, r) in enumerate(amostra.iterrows(), start=1):
        # dois itens propositalmente com preço igual; os demais entre -15 % e +22 %
        fator = 1.00 if i in (4, 13) else rng.choice([0.85, 0.90, 0.94, 0.97, 1.03, 1.06, 1.10, 1.15, 1.22])
        linhas.append({
            "codigo_interno": f"SKU-{i:03d}",
            "titulo": r["titulo"],
            "meu_preco": round(r["preco_concorrente"] * fator, 2),
        })
    # dois itens que NÃO existem no concorrente, para mostrar o tratamento de "sem par"
    linhas.append({"codigo_interno": "SKU-021", "titulo": "Produto exclusivo da minha loja A", "meu_preco": 39.90})
    linhas.append({"codigo_interno": "SKU-022", "titulo": "Produto exclusivo da minha loja B", "meu_preco": 59.90})
    df = pd.DataFrame(linhas)
    ARQ_MEUS_PRECOS.parent.mkdir(exist_ok=True)
    df.to_csv(ARQ_MEUS_PRECOS, index=False, **CSV_OPCOES)
    log.info("gerado %s com %d linhas", ARQ_MEUS_PRECOS.name, len(df))
    return df


def comparar(catalogo: pd.DataFrame, meus: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Devolve (comparativo, sem_par). Chave de ligação: título normalizado."""
    def chave(s: pd.Series) -> pd.Series:
        return s.astype(str).str.strip().str.lower()

    cat = catalogo[["titulo", "preco_concorrente", "avaliacao_estrelas", "disponibilidade", "url"]].copy()
    cat["_k"] = chave(cat["titulo"])
    cat = cat.drop_duplicates("_k")
    m = meus.copy()
    m["_k"] = chave(m["titulo"])

    comp = m.merge(cat.drop(columns=["titulo"]), on="_k", how="left")
    sem_par = comp[comp["preco_concorrente"].isna()].drop(columns=["_k"])
    comp = comp.dropna(subset=["preco_concorrente"]).drop(columns=["_k"]).copy()

    comp["diferenca"] = (comp["meu_preco"] - comp["preco_concorrente"]).round(2)
    comp["diferenca_pct"] = (comp["diferenca"] / comp["preco_concorrente"] * 100).round(2)

    def situacao(d: float) -> str:
        if abs(d) < 0.005:
            return "igual"
        return "estou mais caro" if d > 0 else "estou mais barato"

    comp["situacao"] = comp["diferenca"].apply(situacao)
    comp["alerta"] = comp["diferenca_pct"] > LIMIAR_ALERTA_PCT
    comp = comp.sort_values("diferenca_pct", ascending=False).reset_index(drop=True)
    return comp, sem_par


def calcular_kpis(comp: pd.DataFrame) -> dict:
    return {
        "monitorados": int(len(comp)),
        "mais_caro": int((comp["situacao"] == "estou mais caro").sum()),
        "mais_barato": int((comp["situacao"] == "estou mais barato").sum()),
        "iguais": int((comp["situacao"] == "igual").sum()),
        "alertas": int(comp["alerta"].sum()),
        "dif_media_pct": float(comp["diferenca_pct"].mean()) if len(comp) else 0.0,
    }


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def gerar_xlsx(comp: pd.DataFrame, catalogo: pd.DataFrame, sem_par: pd.DataFrame) -> Path:
    ARQ_XLSX.parent.mkdir(exist_ok=True)

    col_comp = {
        "codigo_interno": "Código", "titulo": "Produto", "meu_preco": "Meu preço",
        "preco_concorrente": "Preço concorrente", "diferenca": "Diferença",
        "diferenca_pct": "Diferença %", "situacao": "Situação",
        "avaliacao_estrelas": "Estrelas", "disponibilidade": "Disponibilidade", "url": "URL",
    }
    aba_comp = comp[list(col_comp)].rename(columns=col_comp)

    col_cat = {
        "data": "Data", "titulo": "Produto", "preco_concorrente": "Preço",
        "avaliacao_estrelas": "Estrelas", "disponibilidade": "Disponibilidade",
        "url": "URL", "coletado_em": "Coletado em",
    }
    aba_cat = catalogo[list(col_cat)].rename(columns=col_cat)

    with pd.ExcelWriter(ARQ_XLSX, engine="openpyxl") as w:
        aba_comp.to_excel(w, sheet_name="Comparativo", index=False)
        aba_cat.to_excel(w, sheet_name="Catalogo_raspado", index=False)
        if len(sem_par):
            sem_par[["codigo_interno", "titulo", "meu_preco"]].rename(
                columns={"codigo_interno": "Código", "titulo": "Produto", "meu_preco": "Meu preço"}
            ).to_excel(w, sheet_name="Sem_par_no_concorrente", index=False)

        negrito = Font(bold=True)
        cab_fill = PatternFill("solid", fgColor="DDE3EA")
        vermelho = PatternFill("solid", fgColor="F8D0D0")
        verde = PatternFill("solid", fgColor="D3EFD8")

        for ws in w.book.worksheets:
            for cel in ws[1]:
                cel.font = negrito
                cel.fill = cab_fill
                cel.alignment = Alignment(vertical="center")
            ws.auto_filter.ref = ws.dimensions
            ws.freeze_panes = "A2"
            # largura das colunas pelo conteúdo (limitado a 60)
            for idx, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 200)), start=1):
                maior = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(idx)].width = min(max(10, maior + 2), 60)

        ws = w.book["Comparativo"]
        cabecalho = [c.value for c in ws[1]]
        ci = {nome: i + 1 for i, nome in enumerate(cabecalho)}
        for linha in range(2, ws.max_row + 1):
            for nome in ("Meu preço", "Preço concorrente", "Diferença"):
                ws.cell(linha, ci[nome]).number_format = "#,##0.00"
            ws.cell(linha, ci["Diferença %"]).number_format = '0.00"%"'
            pct = ws.cell(linha, ci["Diferença %"]).value or 0
            fill = vermelho if pct > LIMIAR_ALERTA_PCT else (verde if pct < 0 else None)
            if fill:
                for c in range(1, ws.max_column + 1):
                    ws.cell(linha, c).fill = fill
            cel_url = ws.cell(linha, ci["URL"])
            if cel_url.value:
                cel_url.hyperlink = cel_url.value
                cel_url.font = Font(color="0563C1", underline="single")

        ws2 = w.book["Catalogo_raspado"]
        for linha in range(2, ws2.max_row + 1):
            ws2.cell(linha, 3).number_format = "#,##0.00"

    log.info("planilha gerada: %s", ARQ_XLSX)
    return ARQ_XLSX


# ---------------------------------------------------------------------------
# HTML autocontido
# ---------------------------------------------------------------------------
def _fmt_moeda(v: float) -> str:
    return f"{MOEDA} {v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _fmt_pct(v: float) -> str:
    s = f"{v:+.1f}".replace(".", ",")
    return s + "%"


def _svg_barras(comp: pd.DataFrame) -> str:
    """Barras horizontais da diferença % por produto: vermelho acima, verde abaixo."""
    if comp.empty:
        return "<p>Sem dados para o gráfico.</p>"
    esc = html_mod.escape
    linha_h, topo = 26, 24
    w_total, x_rot, w_rot = 860, 10, 300
    x_zero = w_rot + (w_total - w_rot) / 2
    escala = ((w_total - w_rot) / 2 - 60) / max(1.0, float(comp["diferenca_pct"].abs().max()))
    altura = topo + linha_h * len(comp) + 10

    partes = [f'<svg viewBox="0 0 {w_total} {altura}" width="100%" role="img" '
              f'aria-label="Diferença percentual do meu preço em relação ao concorrente" '
              f'style="font-family:inherit;font-size:12px">']
    partes.append(f'<line x1="{x_zero}" y1="{topo - 6}" x2="{x_zero}" y2="{altura - 6}" '
                  f'stroke="#9aa4b2" stroke-width="1"/>')
    partes.append(f'<text x="{x_zero}" y="12" text-anchor="middle" fill="#6b7280">0%</text>')
    partes.append(f'<text x="{x_zero - 8}" y="12" text-anchor="end" fill="#2e7d32">mais barato ◀</text>')
    partes.append(f'<text x="{x_zero + 8}" y="12" text-anchor="start" fill="#c62828">▶ mais caro</text>')

    for i, (_, r) in enumerate(comp.iterrows()):
        y = topo + i * linha_h
        pct = float(r["diferenca_pct"])
        titulo = str(r["titulo"])
        rot = titulo if len(titulo) <= 40 else titulo[:38].rstrip() + "…"
        cor = "#c62828" if pct > 0 else ("#2e7d32" if pct < 0 else "#9aa4b2")
        largura = max(2, abs(pct) * escala)
        x = x_zero if pct >= 0 else x_zero - largura
        partes.append(f'<title></title>')
        partes.append(f'<g><title>{esc(titulo)}: {esc(_fmt_pct(pct))}</title>')
        partes.append(f'<text x="{x_rot}" y="{y + 16}" fill="#1f2937">{esc(rot)}</text>')
        partes.append(f'<rect x="{x:.1f}" y="{y + 4}" width="{largura:.1f}" height="{linha_h - 8}" '
                      f'rx="3" fill="{cor}" opacity="0.9"/>')
        tx = x + largura + 6 if pct >= 0 else x - 6
        anc = "start" if pct >= 0 else "end"
        partes.append(f'<text x="{tx:.1f}" y="{y + 16}" text-anchor="{anc}" fill="#374151">'
                      f'{esc(_fmt_pct(pct))}</text></g>')
    partes.append("</svg>")
    return "\n".join(p for p in partes if p != "<title></title>")


def gerar_html(comp: pd.DataFrame, kpis: dict, sem_par: pd.DataFrame, catalogo: pd.DataFrame) -> Path:
    ARQ_HTML.parent.mkdir(exist_ok=True)
    esc = html_mod.escape

    linhas = []
    for _, r in comp.iterrows():
        pct = float(r["diferenca_pct"])
        classe = "caro" if pct > LIMIAR_ALERTA_PCT else ("barato" if pct < 0 else "")
        estrelas = int(r["avaliacao_estrelas"]) if pd.notna(r["avaliacao_estrelas"]) else 0
        linhas.append(
            f'<tr class="{classe}">'
            f'<td>{esc(str(r["codigo_interno"]))}</td>'
            f'<td><a href="{esc(str(r["url"]))}" target="_blank" rel="noopener">{esc(str(r["titulo"]))}</a></td>'
            f'<td class="num" data-v="{r["meu_preco"]}">{_fmt_moeda(r["meu_preco"])}</td>'
            f'<td class="num" data-v="{r["preco_concorrente"]}">{_fmt_moeda(r["preco_concorrente"])}</td>'
            f'<td class="num" data-v="{r["diferenca"]}">{_fmt_moeda(r["diferenca"])}</td>'
            f'<td class="num" data-v="{pct}"><b>{_fmt_pct(pct)}</b></td>'
            f'<td>{esc(str(r["situacao"]))}</td>'
            f'<td class="num" data-v="{estrelas}" title="{estrelas} de 5">{"★" * estrelas}{"☆" * (5 - estrelas)}</td>'
            f'<td>{esc(str(r["disponibilidade"]))}</td>'
            f'</tr>'
        )

    sem_par_html = ""
    if len(sem_par):
        itens = "".join(f"<li>{esc(str(r['codigo_interno']))} — {esc(str(r['titulo']))}</li>"
                        for _, r in sem_par.iterrows())
        sem_par_html = (f'<section class="card"><h2>Itens da minha tabela sem par no concorrente '
                        f'({len(sem_par)})</h2><ul>{itens}</ul></section>')

    media = kpis["dif_media_pct"]
    cor_media = "#c62828" if media > 0 else "#2e7d32"
    gerado_em = datetime.now().strftime("%d/%m/%Y %H:%M")

    pagina = f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Monitor de preços — demo</title>
<style>
  :root {{ --bg:#f4f6f8; --card:#fff; --txt:#1f2937; --mut:#6b7280; --borda:#e5e7eb;
          --verm:#c62828; --verd:#2e7d32; --verm-bg:#fdecec; --verd-bg:#e8f5e9; }}
  * {{ box-sizing:border-box; }}
  body {{ margin:0; background:var(--bg); color:var(--txt);
         font:15px/1.45 system-ui,-apple-system,"Segoe UI",Roboto,sans-serif; }}
  .wrap {{ max-width:1100px; margin:0 auto; padding:20px 16px 40px; }}
  header h1 {{ margin:0 0 4px; font-size:1.6rem; }}
  header p {{ margin:0 0 18px; color:var(--mut); }}
  .kpis {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(150px,1fr)); gap:12px; margin-bottom:18px; }}
  .kpi {{ background:var(--card); border:1px solid var(--borda); border-radius:10px; padding:14px 16px; }}
  .kpi .v {{ font-size:1.9rem; font-weight:700; line-height:1.1; }}
  .kpi .l {{ color:var(--mut); font-size:.85rem; margin-top:4px; }}
  .card {{ background:var(--card); border:1px solid var(--borda); border-radius:10px; padding:16px; margin-bottom:18px; }}
  .card h2 {{ margin:0 0 10px; font-size:1.1rem; }}
  .card p.ajuda {{ color:var(--mut); font-size:.85rem; margin:0 0 8px; }}
  .tabela-wrap {{ overflow-x:auto; }}
  table {{ border-collapse:collapse; width:100%; min-width:760px; font-size:.92rem; }}
  th, td {{ padding:8px 10px; border-bottom:1px solid var(--borda); text-align:left; vertical-align:top; }}
  th {{ background:#f9fafb; cursor:pointer; user-select:none; white-space:nowrap; position:sticky; top:0; }}
  th:hover {{ background:#eef2f6; }}
  th .seta {{ color:var(--mut); font-size:.75rem; margin-left:4px; }}
  td.num {{ text-align:right; white-space:nowrap; font-variant-numeric:tabular-nums; }}
  tr.caro td {{ background:var(--verm-bg); }}
  tr.barato td {{ background:var(--verd-bg); }}
  a {{ color:#0b5cad; text-decoration:none; }} a:hover {{ text-decoration:underline; }}
  footer {{ color:var(--mut); font-size:.85rem; text-align:center; margin-top:24px; }}
  @media (max-width:600px) {{ .kpi .v {{ font-size:1.5rem; }} header h1 {{ font-size:1.3rem; }} }}
</style>
</head>
<body>
<div class="wrap">
<header>
  <h1>Monitor de preços — demo</h1>
  <p>Comparação entre a minha tabela e o catálogo do concorrente · coleta de {esc(gerado_em)} ·
     {len(catalogo)} produtos raspados em {PAGINAS} páginas</p>
</header>

<section class="kpis" id="kpis">
  <div class="kpi"><div class="v">{kpis["monitorados"]}</div><div class="l">produtos monitorados</div></div>
  <div class="kpi"><div class="v" style="color:var(--verm)">{kpis["mais_caro"]}</div><div class="l">estou mais caro</div></div>
  <div class="kpi"><div class="v" style="color:var(--verd)">{kpis["mais_barato"]}</div><div class="l">estou mais barato</div></div>
  <div class="kpi"><div class="v" style="color:{cor_media}">{_fmt_pct(media)}</div><div class="l">diferença média</div></div>
  <div class="kpi"><div class="v">{kpis["alertas"]}</div><div class="l">alertas (&gt; {LIMIAR_ALERTA_PCT:.0f}% acima)</div></div>
</section>

<section class="card">
  <h2>Diferença % por produto</h2>
  <p class="ajuda">Meu preço em relação ao concorrente. Barra à direita (vermelho) = estou mais caro; à esquerda (verde) = mais barato.</p>
  {_svg_barras(comp)}
</section>

<section class="card">
  <h2>Comparativo</h2>
  <p class="ajuda">Clique no cabeçalho para ordenar. Linhas vermelhas: mais de {LIMIAR_ALERTA_PCT:.0f}% acima do concorrente; verdes: abaixo.</p>
  <div class="tabela-wrap">
  <table id="tab">
    <thead><tr>
      <th>Código</th><th>Produto</th><th class="num">Meu preço</th><th class="num">Concorrente</th>
      <th class="num">Diferença</th><th class="num">Dif. %</th><th>Situação</th><th class="num">Estrelas</th><th>Disponib.</th>
    </tr></thead>
    <tbody>
    {"".join(linhas)}
    </tbody>
  </table>
  </div>
</section>

{sem_par_html}

<footer>Demo com dados reais de um site de treino de scraping (books.toscrape.com) e preços próprios fictícios · Marcos Galani · automação de processos</footer>
</div>

<script>
(function () {{
  var tab = document.getElementById('tab');
  var ths = tab.querySelectorAll('th');
  var dir = {{}};
  ths.forEach(function (th, idx) {{
    th.addEventListener('click', function () {{
      var asc = dir[idx] = !dir[idx];
      ths.forEach(function (t) {{ var s = t.querySelector('.seta'); if (s) s.remove(); }});
      var seta = document.createElement('span'); seta.className = 'seta'; seta.textContent = asc ? '▲' : '▼';
      th.appendChild(seta);
      var tb = tab.tBodies[0];
      var linhas = Array.prototype.slice.call(tb.rows);
      linhas.sort(function (a, b) {{
        var ca = a.cells[idx], cb = b.cells[idx];
        var va = ca.dataset.v !== undefined ? parseFloat(ca.dataset.v) : ca.textContent.trim().toLowerCase();
        var vb = cb.dataset.v !== undefined ? parseFloat(cb.dataset.v) : cb.textContent.trim().toLowerCase();
        if (va < vb) return asc ? -1 : 1;
        if (va > vb) return asc ? 1 : -1;
        return 0;
      }});
      linhas.forEach(function (l) {{ tb.appendChild(l); }});
    }});
  }});
}})();
</script>
</body>
</html>
"""
    ARQ_HTML.write_text(pagina, encoding="utf-8")
    log.info("página gerada: %s", ARQ_HTML)
    return ARQ_HTML


# ---------------------------------------------------------------------------
def main() -> int:
    log.info("== Monitor de preços — início ==")
    catalogo = raspar()
    if catalogo.empty:
        log.error("nenhum produto raspado; confira URL_PAGINA e os seletores")
        return 1
    salvar_historico(catalogo)
    meus = garantir_meus_precos(catalogo)
    comp, sem_par = comparar(catalogo, meus)
    kpis = calcular_kpis(comp)
    gerar_xlsx(comp, catalogo, sem_par)
    gerar_html(comp, kpis, sem_par, catalogo)

    print()
    print(f"Produtos raspados ........ {len(catalogo)}")
    print(f"Produtos monitorados ..... {kpis['monitorados']}  (sem par no concorrente: {len(sem_par)})")
    print(f"Estou mais caro .......... {kpis['mais_caro']}  (alertas > {LIMIAR_ALERTA_PCT:.0f}%: {kpis['alertas']})")
    print(f"Estou mais barato ........ {kpis['mais_barato']}")
    print(f"Iguais ................... {kpis['iguais']}")
    print(f"Diferença média .......... {kpis['dif_media_pct']:+.2f}%")
    print(f"Planilha ................. {ARQ_XLSX}")
    print(f"Página ................... {ARQ_HTML}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
